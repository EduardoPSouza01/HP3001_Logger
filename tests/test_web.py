import io

from conftest import seed
from openpyxl import load_workbook


def test_pages_render(client):
    c, repo = client
    seed(repo)
    assert c.get("/").status_code == 200
    r = c.get("/history?channel=1")
    assert r.status_code == 200 and b"Camara" in r.data
    assert c.get("/history?start=lixo&page=abc").status_code == 200


def test_api(client):
    c, repo = client
    seed(repo)
    latest = c.get("/api/latest").get_json()
    assert [i["channel"] for i in latest] == [1, 5]
    assert latest[0]["name"] == "Camara" and latest[0]["stats_24h"]["temp_max"] == 22.0
    series = c.get("/api/series?hours=1").get_json()
    assert len(series) == 2 and len(series[0]["points"]) == 3
    assert c.get("/api/status").get_json()["running"] is False


def test_export_xlsx(client):
    c, repo = client
    seed(repo)
    r = c.get("/export.xlsx?channel=5")
    assert r.status_code == 200 and "attachment" in r.headers["Content-Disposition"]
    ws = load_workbook(io.BytesIO(r.data)).active
    assert ws.max_row == 4 and ws["C2"].value == 5
